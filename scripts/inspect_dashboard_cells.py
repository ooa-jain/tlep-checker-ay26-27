import sys
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

wb = openpyxl.load_workbook('OOA_TLEP_Review_Checklist_AY_2026-27.xlsx', data_only=False)
ws = wb['Dashboard']
for r in range(1, ws.max_row+1):
    vals = []
    for c in range(1, ws.max_column+1):
        v = ws.cell(r, c).value
        if v is not None:
            vals.append(f"{openpyxl.utils.get_column_letter(c)}{r}: {v}")
    if vals:
        print(" | ".join(vals))
