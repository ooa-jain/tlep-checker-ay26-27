import sys
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl, docx

print("=== Checking TLEP 2023-24,24-25 and 25-26.xlsx ===")
try:
    wb = openpyxl.load_workbook(r'C:\Users\jljga\Downloads\TLEP 2023-24,24-25 and 25-26.xlsx', data_only=True)
    print("Sheets:", wb.sheetnames)
    for name in wb.sheetnames:
        ws = wb[name]
        print(f"Sheet {name}: max_row={ws.max_row}, max_col={ws.max_col}")
        for r in range(1, min(ws.max_row+1, 15)):
            row_vals = [str(ws.cell(r, c).value) if ws.cell(r, c).value is not None else '' for c in range(1, min(ws.max_col+1, 10))]
            if any(row_vals):
                print(f"  R{r}:", row_vals[:6])
except Exception as e:
    print("Error:", e)

print("\n=== Checking tlep writeup.docx ===")
try:
    doc = docx.Document(r'C:\Users\jljga\Downloads\tlep writeup.docx')
    print("Paragraphs count:", len(doc.paragraphs))
    for p in doc.paragraphs[:10]:
        if p.text.strip():
            print("  P:", p.text.strip())
    print("Tables count:", len(doc.tables))
    for i, table in enumerate(doc.tables[:3]):
        print(f"  Table {i} rows={len(table.rows)} cols={len(table.columns)}")
        for row in table.rows[:3]:
            print("    Row:", [cell.text.strip() for cell in row.cells[:5]])
except Exception as e:
    print("Error:", e)
