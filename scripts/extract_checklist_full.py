import sys
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl, json, os

wb = openpyxl.load_workbook('OOA_TLEP_Review_Checklist_AY_2026-27.xlsx', data_only=False)

# 1. Sheet names
sheets = wb.sheetnames
print("Sheets:", sheets)

# 2. Extract Checklist rows
ws_check = wb['TLEP Review Checklist']
params = []

# Headers are at Row 4
headers = [ws_check.cell(4, c).value for c in range(1, 12)]
print("Checklist Headers:", headers)

for r in range(5, 54):
    p_id = ws_check.cell(r, 1).value
    area = ws_check.cell(r, 2).value
    param = ws_check.cell(r, 3).value
    criteria = ws_check.cell(r, 4).value
    priority = ws_check.cell(r, 5).value
    default_status = ws_check.cell(r, 6).value
    score_formula = ws_check.cell(r, 7).value
    remarks = ws_check.cell(r, 8).value
    action_req = ws_check.cell(r, 9).value
    reviewed_by = ws_check.cell(r, 10).value
    review_date = ws_check.cell(r, 11).value
    
    # Classify validation type based on nature
    # Deterministic: structural, numerical, exact match
    # Hybrid: structural checks + AI semantic judgment
    # AI: qualitative academic judgment
    deterministic_ids = {1, 2, 3, 4, 5, 6, 7, 8, 14, 19, 20, 21, 24, 27, 29, 30, 31, 32, 36, 37, 38, 40}
    ai_ids = {9, 10, 11, 15, 26, 33, 34, 35, 43, 44, 45}
    # Others are hybrid (12, 13, 16, 17, 18, 22, 23, 25, 28, 39, 41, 42, 46, 47, 48, 49)
    
    p_int = int(p_id)
    if p_int in deterministic_ids:
        val_type = "deterministic"
    elif p_int in ai_ids:
        val_type = "ai"
    else:
        val_type = "hybrid"

    params.append({
        "parameter_id": p_int,
        "review_area": str(area).strip(),
        "parameter": str(param).strip(),
        "criterion": str(criteria).strip(),
        "priority": str(priority).strip(),
        "validation_type": val_type,
        "default_status": str(default_status).strip() if default_status else "Not Reviewed",
        "score_formula": str(score_formula).strip() if score_formula else "",
        "row_index": r
    })

print(f"Total parameters extracted: {len(params)}")

# 3. Instructions & Scoring Methodology
ws_inst = wb['Instructions']
instructions = []
for r in range(1, ws_inst.max_row + 1):
    vals = [ws_inst.cell(r, c).value for c in range(1, ws_inst.max_column + 1)]
    vals = [str(v).strip() for v in vals if v is not None and str(v).strip()]
    if vals:
        instructions.append({"row": r, "content": " : ".join(vals)})

# 4. Hours Validation structure
ws_hours = wb['Hours Validation']
hours_items = []
for r in range(4, 16):
    param_name = ws_hours.cell(r, 1).value
    approved = ws_hours.cell(r, 2).value
    tlep = ws_hours.cell(r, 3).value
    var_formula = ws_hours.cell(r, 4).value
    status = ws_hours.cell(r, 5).value
    rem = ws_hours.cell(r, 6).value
    hours_items.append({
        "row": r,
        "parameter": str(param_name).strip() if param_name else "",
        "variance_formula": str(var_formula).strip() if var_formula else ""
    })

# 5. CO-PO Review structure
ws_copo = wb['CO-PO Review']
copo_headers = [ws_copo.cell(3, c).value for c in range(1, 9)]

checklist_metadata = {
    "title": "OOA TLEP Academic Review Checklist | AY 2026–27",
    "version": "AY 2026–27",
    "source_file": "OOA_TLEP_Review_Checklist_AY_2026-27.xlsx",
    "total_parameters": len(params),
    "scoring_system": {
        "Compliant": 2,
        "Needs Revision": 1,
        "Major Revision": 0,
        "Non-Compliant": 0,
        "NA": "Excluded from numerator and denominator",
        "Not Reviewed": "Excluded from numerator and denominator"
    },
    "score_calculation": {
        "applicable_parameters": "Total Parameters - NA Parameters - Not Reviewed Parameters",
        "max_score": "Applicable Parameters * 2",
        "compliance_percentage": "(Obtained Score / Maximum Score) * 100"
    },
    "sheets": sheets,
    "instructions": instructions,
    "hours_validation_parameters": hours_items,
    "copo_review_headers": [str(h) for h in copo_headers if h is not None],
    "parameters": params
}

os.makedirs('config', exist_ok=True)
with open('config/checklist.json', 'w', encoding='utf-8') as f:
    json.dump(checklist_metadata, f, indent=2, ensure_ascii=False)

print("Saved config/checklist.json successfully with", len(params), "parameters.")
