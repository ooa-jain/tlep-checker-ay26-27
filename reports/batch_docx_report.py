import io
import pandas as pd
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def is_pg(programme: str, course_title: str) -> bool:
    prog_lower = str(programme).lower()
    title_lower = str(course_title).lower()
    
    pg_keywords = ["m.", "master", "mba", "mca", "m.tech", "m.sc", "m.com", "m.des", "pg", "post graduate"]
    ug_keywords = ["b.", "bachelor", "bba", "bca", "b.tech", "b.sc", "b.com", "b.des", "ug", "under graduate"]
    
    for kw in pg_keywords:
        if kw in prog_lower or kw in title_lower:
            return True
    
    for kw in ug_keywords:
        if kw in prog_lower or kw in title_lower:
            return False
            
    return False

def generate_programme_executive_docx(results) -> bytes:
    doc = Document()
    
    # Title
    title = doc.add_paragraph("INSTITUTIONAL TLEP PORTFOLIO – PROGRAMME EXECUTIVE REPORT")
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    for run in title.runs:
        run.bold = True
        run.font.size = Pt(14)
        
    doc.add_paragraph()
    
    if not results:
        doc.add_paragraph("No audited courses available.")
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()
        
    data = []
    for r in results:
        level = "PG" if is_pg(getattr(r.normalized_tlep.course_info, 'programme', ''), getattr(r.normalized_tlep.course_info, 'course_title', '')) else "UG"
        data.append({
            "Score": r.compliance_percentage,
            "Status": r.overall_status.value,
            "Level": level,
            "Department": getattr(r.normalized_tlep.course_info, 'department', 'General')
        })
        
    df = pd.DataFrame(data)
    
    total_courses = len(df)
    avg_score = df["Score"].mean()
    compliant = len(df[df["Status"] == "Compliant"])
    needs_rev = len(df[df["Status"] == "Needs Revision"])
    major_rev = len(df[df["Status"].isin(["Major Revision", "Non-Compliant"])])
    
    # 1. Status Distribution Chart (Top)
    plt.figure(figsize=(6, 4))
    status_counts = df["Status"].value_counts()
    colors = {"Compliant": "#16A34A", "Needs Revision": "#CA8A04", "Major Revision": "#DC2626", "Non-Compliant": "#991B1B"}
    plot_colors = [colors.get(s, "#888888") for s in status_counts.index]
    
    wedges, texts, autotexts = plt.pie(
        status_counts.values, 
        labels=status_counts.index, 
        autopct='%1.1f%%', 
        startangle=140, 
        colors=plot_colors,
        wedgeprops=dict(width=0.45, edgecolor='white', linewidth=2),
        textprops=dict(color='#1E293B', fontsize=10, fontweight='500')
    )
    for autotext in autotexts:
        autotext.set_color('white')
        autotext.set_weight('bold')
        
    plt.title("Compliance Status Distribution", pad=15, fontsize=12, fontweight='bold', color='#0F172A')
    buf1 = io.BytesIO()
    plt.savefig(buf1, format='png', bbox_inches='tight', dpi=300, transparent=True)
    buf1.seek(0)
    plt.close()
    
    doc.add_paragraph("1. Compliance Status Distribution").runs[0].bold = True
    doc.add_picture(buf1, width=Inches(4.5))
    doc.add_paragraph()

    # Summary Statistics text
    stats_p = doc.add_paragraph()
    stats_p.add_run("PROGRAMME STATISTICS SUMMARY\n").bold = True
    stats_p.add_run(f"Total Courses Audited: {total_courses}\n")
    stats_p.add_run(f"Average Compliance Score: {avg_score:.1f}%\n")
    stats_p.add_run(f"Fully Compliant (Ready for BoS): {compliant}\n")
    stats_p.add_run(f"Needs Minor Revision: {needs_rev}\n")
    stats_p.add_run(f"Requires Major Rework: {major_rev}\n")
    doc.add_paragraph()
    
    # 2. UG vs PG Breakdown Chart
    plt.figure(figsize=(5, 4))
    level_counts = df["Level"].value_counts()
    plt.bar(level_counts.index, level_counts.values, color=["#2563EB", "#7C3AED"])
    plt.title("UG vs PG Course Distribution")
    plt.ylabel("Number of Courses")
    buf2 = io.BytesIO()
    plt.savefig(buf2, format='png', bbox_inches='tight')
    buf2.seek(0)
    plt.close()
    
    doc.add_paragraph("2. UG vs PG Course Breakdown").runs[0].bold = True
    doc.add_picture(buf2, width=Inches(4.0))
    
    # 3. Department Performance Chart
    dept_scores = df.groupby("Department")["Score"].mean().sort_values(ascending=True)
    plt.figure(figsize=(6, max(3, len(dept_scores) * 0.4)))
    plt.barh(dept_scores.index, dept_scores.values, color="#0EA5E9")
    plt.title("Average Compliance Score by Department")
    plt.xlabel("Average Compliance (%)")
    plt.xlim(0, 100)
    buf3 = io.BytesIO()
    plt.savefig(buf3, format='png', bbox_inches='tight')
    buf3.seek(0)
    plt.close()
    
    doc.add_paragraph("3. Department Performance Leaderboard").runs[0].bold = True
    doc.add_picture(buf3, width=Inches(5.5))
    
    buf_out = io.BytesIO()
    doc.save(buf_out)
    return buf_out.getvalue()
