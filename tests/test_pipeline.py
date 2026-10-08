"""
End-to-End Pipeline Test for OOA TLEP Compliance Review System
AY 2026–27
"""

import sys
import os
sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from engine.reviewer import review_tlep_document
from reports.excel_report import generate_excel_report
from models.schemas import StatusEnum

def create_sample_test_tlep(file_path: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Course Details & Syllabus"
    
    # 1. Course Info
    ws.append(["Course Information", ""])
    ws.append(["Course Title", "Machine Learning Systems"])
    ws.append(["Course Code", "CS601"])
    ws.append(["Semester", "Semester VI"])
    ws.append(["Academic Year", "AY 2026–27"])
    ws.append(["Credits", 4.0])
    ws.append(["L-T-P-E", "3-0-2-0"])
    ws.append(["CA : ESE", "50:50"])
    ws.append(["Pass Marks", "Minimum 40% in CA and 40% in ESE"])
    ws.append([])
    
    # 2. Course Outcomes
    ws.append(["Course Outcomes", "Statement", "BTL"])
    ws.append(["CO1", "Explain the fundamental principles of supervised learning algorithms", "Level 2 - Understand"])
    ws.append(["CO2", "Implement gradient descent optimization algorithms for linear models", "Level 3 - Apply"])
    ws.append(["CO3", "Analyze model performance using precision, recall, and ROC curves", "Level 4 - Analyze"])
    ws.append(["CO4", "Design deep neural network architectures for computer vision tasks", "Level 6 - Create"])
    ws.append([])
    
    # 3. CO-PO Matrix
    ws.append(["CO", "PO1", "PO2", "PO3", "PO4", "PO5", "PSO1", "PSO2"])
    ws.append(["CO1", "3", "2", "-", "-", "1", "3", "2"])
    ws.append(["CO2", "3", "3", "2", "-", "2", "3", "3"])
    ws.append(["CO3", "3", "3", "3", "2", "2", "3", "2"])
    ws.append(["CO4", "3", "3", "3", "3", "3", "3", "3"])
    ws.append([])
    
    # 4. Session Plan
    ws_sess = wb.create_sheet(title="Session Plan")
    ws_sess.append(["Session No", "Module", "Topic", "Pedagogy", "Mode", "CO", "Reading"])
    for s in range(1, 31):
        mod_idx = f"Module {((s - 1) // 8) + 1}"
        co_idx = f"CO{((s - 1) // 8) + 1}"
        ws_sess.append([
            s,
            mod_idx,
            f"Topic {s}: Core Algorithmic Concepts in {mod_idx}",
            "Interactive Lecture & Coding Demo",
            "L - Synchronous",
            co_idx,
            "Goodfellow Ch. 5"
        ])
    
    # 5. Assessments
    ws_ass = wb.create_sheet(title="Assessments")
    ws_ass.append(["Assessment Component", "Weightage", "Type", "CO Mapped", "BTL"])
    ws_ass.append(["Continuous Internal Assessment 1", "20%", "Formative", "CO1, CO2", "K3"])
    ws_ass.append(["Continuous Internal Assessment 2", "20%", "Formative", "CO3", "K4"])
    ws_ass.append(["Lab Mini Project / Assignment", "10%", "Formative", "CO4", "K6"])
    ws_ass.append(["End Semester Examination", "50%", "Summative", "CO1, CO2, CO3, CO4", "K4"])
    
    # 6. Textbooks
    ws_books = wb.create_sheet(title="Learning Resources")
    ws_books.append(["Category", "Title & Author", "Publisher / Year"])
    ws_books.append(["Textbook", "Pattern Recognition and Machine Learning, Christopher Bishop", "Springer, 2006"])
    ws_books.append(["Reference", "Deep Learning, Ian Goodfellow, Yoshua Bengio", "MIT Press, 2016"])
    ws_books.append(["Web Resource", "Coursera Deep Learning Specialization", "deeplearning.ai"])
    
    wb.save(file_path)
    print(f"Created sample TLEP workbook at: {file_path}")


def main():
    test_dir = "tests/test_data"
    os.makedirs(test_dir, exist_ok=True)
    sample_file = os.path.join(test_dir, "Sample_TLEP_CS601.xlsx")
    create_sample_test_tlep(sample_file)
    
    print("\nRunning compliance review on Sample_TLEP_CS601.xlsx...")
    result = review_tlep_document(sample_file)
    
    print("\n================ AUDIT SUMMARY ================")
    print(f"Review ID: {result.review_id}")
    print(f"File Name: {result.file_name}")
    print(f"Total Parameters: {result.total_parameters}")
    print(f"Compliant: {result.compliant_count}")
    print(f"Needs Revision: {result.needs_revision_count}")
    print(f"Major Revision: {result.major_revision_count}")
    print(f"Non-Compliant: {result.non_compliant_count}")
    print(f"NA: {result.na_count}")
    print(f"Score Obtained: {result.score_obtained} / {result.maximum_score}")
    print(f"Compliance %: {result.compliance_percentage}%")
    print(f"Overall Status: {result.overall_status.value}")
    
    assert len(result.parameter_findings) == 49, f"Expected 49 findings, got {len(result.parameter_findings)}"
    assert result.total_parameters == 49
    assert len(result.hours_validation_rows) == 12, f"Expected 12 hours validation rows, got {len(result.hours_validation_rows)}"
    
    # Test Excel report export
    report_file = os.path.join(test_dir, "Sample_TLEP_Audit_Report.xlsx")
    generate_excel_report(result, output_path=report_file)
    assert os.path.exists(report_file), "Excel report file was not created!"
    print(f"\nGenerated downloadable audit report at: {report_file}")
    print("All assertions passed successfully!")

if __name__ == "__main__":
    main()
